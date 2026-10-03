from __future__ import annotations

import hashlib
import hmac
import logging
import os
import re
import secrets
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal

import psycopg
import uvicorn
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Request, Response
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, field_validator
from psycopg.rows import dict_row

ROOT = Path(__file__).resolve().parent
load_dotenv(ROOT / ".env")
DATABASE_URL = os.getenv("DATABASE_URL")
SESSION_COOKIE = "student_portal_session"
SESSION_MAX_AGE = 30 * 24 * 60 * 60
PASSWORD_HASH_N = 2**14
EXPORT_DIR = ROOT / "exports"
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("student-profile")


def to_camel(name: str) -> str:
    first, *rest = name.split("_")
    return first + "".join(part.capitalize() for part in rest)


class FormModel(BaseModel):
    model_config = ConfigDict(alias_generator=to_camel, populate_by_name=True, extra="forbid")

    @field_validator("*", mode="before")
    @classmethod
    def empty_text_is_missing(cls, value: object) -> object:
        return None if value == "" else value


class Credentials(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=12, max_length=128)

    @field_validator("email")
    @classmethod
    def normalize_email(cls, value: str) -> str:
        normalized = value.strip().lower()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
            raise ValueError("Enter a valid email address.")
        return normalized


class Certification(FormModel):
    name: str | None = None
    organization: str | None = None
    year: int | None = None


class Achievement(FormModel):
    title: str | None = None
    type: str | None = None
    year: int | None = None
    description: str | None = None


class ProfileSubmission(FormModel):
    full_name: str | None = None
    preferred_name: str | None = None
    age: int | None = Field(default=None, ge=13, le=99)
    date_of_birth: date | None = None
    gender: str | None = None
    email: str | None = None
    phone: str | None = None
    city: str | None = None

    school: str | None = None
    course: str | None = None
    department: str | None = None
    year: str | None = None
    semester: str | None = None
    student_id: str | None = None
    grade: str | None = None

    favorite_subjects: list[str] = Field(default_factory=list)
    difficult_subject: str | None = None
    studying_subjects: list[str] = Field(default_factory=list)
    study_hours: str | None = None
    learning_methods: list[str] = Field(default_factory=list)
    academic_confidence: int | None = Field(default=None, ge=1, le=5)

    languages: list[str] = Field(default_factory=list)
    technical_skills: list[str] = Field(default_factory=list)
    strongest_skill: str | None = None
    next_technology: str | None = None
    technical_level: str | None = None

    clubs: list[str] = Field(default_factory=list)
    hackathons: Literal["Yes", "No"] | None = None
    workshops: Literal["Yes", "No"] | None = None
    competitions: Literal["Yes", "No"] | None = None
    certifications: list[Certification] = Field(default_factory=list)
    achievements: list[Achievement] = Field(default_factory=list)

    hobbies: list[str] = Field(default_factory=list)
    interest_area: str | None = None
    career_fields: list[str] = Field(default_factory=list)

    career: str | None = None
    goal_skill: str | None = None
    two_year_goal: str | None = None
    organization: str | None = None
    self_improvement: str | None = None


def get_database_url() -> str:
    if not DATABASE_URL:
        raise HTTPException(
            status_code=503,
            detail="Database is not configured. Set DATABASE_URL in .env and run schema.sql in pgAdmin.",
        )
    return DATABASE_URL


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode(), salt=salt, n=PASSWORD_HASH_N, r=8, p=1, dklen=64)
    return f"scrypt${PASSWORD_HASH_N}${salt.hex()}${digest.hex()}"


def verify_password(password: str, stored_hash: str) -> bool:
    try:
        algorithm, n_value, salt_hex, hash_hex = stored_hash.split("$", maxsplit=3)
        n = int(n_value)
        if algorithm != "scrypt" or n < 2**12 or n > 2**18 or n & (n - 1):
            return False
        expected = bytes.fromhex(hash_hex)
        actual = hashlib.scrypt(password.encode(), salt=bytes.fromhex(salt_hex), n=n, r=8, p=1, dklen=len(expected))
    except (ValueError, TypeError):
        return False
    return hmac.compare_digest(actual, expected)


def create_session(cursor: psycopg.Cursor, account_id: int) -> str:
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    cursor.execute("DELETE FROM auth_sessions WHERE expires_at <= NOW()")
    cursor.execute(
        "INSERT INTO auth_sessions (token_hash, account_id, expires_at) "
        "VALUES (%s, %s, NOW() + INTERVAL '30 days')",
        (token_hash, account_id),
    )
    return token


def set_session_cookie(response: Response, token: str) -> None:
    response.set_cookie(
        SESSION_COOKIE,
        token,
        max_age=SESSION_MAX_AGE,
        httponly=True,
        secure=os.getenv("COOKIE_SECURE", "false").lower() == "true",
        samesite="lax",
        path="/",
    )


def authenticated_account(cursor: psycopg.Cursor, request: Request) -> dict[str, object]:
    token = request.cookies.get(SESSION_COOKIE)
    if not token:
        raise HTTPException(status_code=401, detail="Please log in to continue.")
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    cursor.execute(
        """
        SELECT a.id, a.email
        FROM auth_sessions AS session
        JOIN accounts AS a ON a.id = session.account_id
        WHERE session.token_hash = %s AND session.expires_at > NOW()
        """,
        (token_hash,),
    )
    account = cursor.fetchone()
    if account is None:
        raise HTTPException(status_code=401, detail="Your session expired. Please log in again.")
    return account


def read_student_profile(cursor: psycopg.Cursor, account_id: int) -> tuple[int | None, dict[str, object] | None]:
    cursor.execute(
        """
        SELECT
            s.id AS profile_id,
            s.full_name AS "fullName", s.preferred_name AS "preferredName", s.age,
            s.date_of_birth AS "dateOfBirth", s.gender, s.email, s.phone, s.city,
            e.school, e.course, e.department, e.current_year AS year,
            e.current_semester AS semester, e.campus_student_id AS "studentId", e.grade,
            COALESCE(ac.favorite_subjects, ARRAY[]::TEXT[]) AS "favoriteSubjects",
            ac.difficult_subject AS "difficultSubject",
            COALESCE(ac.studying_subjects, ARRAY[]::TEXT[]) AS "studyingSubjects",
            ac.study_hours AS "studyHours",
            COALESCE(ac.learning_methods, ARRAY[]::TEXT[]) AS "learningMethods",
            ac.academic_confidence AS "academicConfidence",
            COALESCE(tech.programming_languages, ARRAY[]::TEXT[]) AS languages,
            COALESCE(tech.technical_skills, ARRAY[]::TEXT[]) AS "technicalSkills",
            tech.strongest_skill AS "strongestSkill", tech.next_technology AS "nextTechnology",
            tech.skill_level AS "technicalLevel",
            COALESCE(act.clubs, ARRAY[]::TEXT[]) AS clubs,
            CASE act.hackathons WHEN TRUE THEN 'Yes' WHEN FALSE THEN 'No' END AS hackathons,
            CASE act.workshops WHEN TRUE THEN 'Yes' WHEN FALSE THEN 'No' END AS workshops,
            CASE act.competitions WHEN TRUE THEN 'Yes' WHEN FALSE THEN 'No' END AS competitions,
            COALESCE(i.hobbies, ARRAY[]::TEXT[]) AS hobbies,
            i.interest_area AS "interestArea",
            COALESCE(i.career_fields, ARRAY[]::TEXT[]) AS "careerFields",
            g.career, g.goal_skill AS "goalSkill", g.two_year_goal AS "twoYearGoal",
            g.organization, g.self_improvement AS "selfImprovement"
        FROM students AS s
        LEFT JOIN education AS e ON e.student_id = s.id
        LEFT JOIN academic_profile AS ac ON ac.student_id = s.id
        LEFT JOIN technical_profile AS tech ON tech.student_id = s.id
        LEFT JOIN activities AS act ON act.student_id = s.id
        LEFT JOIN interests AS i ON i.student_id = s.id
        LEFT JOIN future_goals AS g ON g.student_id = s.id
        WHERE s.account_id = %s
        """,
        (account_id,),
    )
    row = cursor.fetchone()
    if row is None:
        return None, None

    profile = dict(row)
    profile_id = profile.pop("profile_id")
    cursor.execute(
        "SELECT name, organization, year FROM certifications WHERE student_id = %s ORDER BY id",
        (profile_id,),
    )
    profile["certifications"] = [dict(item) for item in cursor.fetchall()]
    cursor.execute(
        "SELECT title, type, year, description FROM achievements WHERE student_id = %s ORDER BY id",
        (profile_id,),
    )
    profile["achievements"] = [dict(item) for item in cursor.fetchall()]
    return profile_id, profile


def excel_value(value: object) -> object:
    if isinstance(value, (list, tuple)):
        value = "; ".join(str(item) for item in value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        value = value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value, str) and value.lstrip(" \t\r").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def style_worksheet(worksheet: object) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="285C4A")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cells in worksheet.columns:
        column_letter = cells[0].column_letter
        width = max((len(str(cell.value or "")) for cell in cells), default=10)
        worksheet.column_dimensions[column_letter].width = min(max(width + 2, 12), 40)


def save_workbook_atomically(workbook: Workbook, target: Path, temporary: Path) -> None:
    output = BytesIO()
    workbook.save(output)
    temporary.write_bytes(output.getvalue())
    try:
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


def build_profile_workbook(profile_id: int, profile: dict[str, object]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    layouts = {
        "Personal": [
            ("Full name", "fullName"), ("Preferred name", "preferredName"), ("Age", "age"),
            ("Date of birth", "dateOfBirth"), ("Gender", "gender"), ("Email", "email"),
            ("Phone", "phone"), ("City", "city"),
        ],
        "Education": [
            ("College / school", "school"), ("Course", "course"), ("Department", "department"),
            ("Year", "year"), ("Semester", "semester"), ("Student ID", "studentId"), ("Grade", "grade"),
        ],
        "Academics": [
            ("Favorite subjects", "favoriteSubjects"), ("Difficult subject", "difficultSubject"),
            ("Studying subjects", "studyingSubjects"), ("Study hours", "studyHours"),
            ("Learning methods", "learningMethods"), ("Academic confidence", "academicConfidence"),
        ],
        "Technical Skills": [
            ("Programming languages", "languages"), ("Technical skills", "technicalSkills"),
            ("Strongest skill", "strongestSkill"), ("Next technology", "nextTechnology"),
            ("Skill level", "technicalLevel"),
        ],
        "Activities": [
            ("Clubs", "clubs"), ("Hackathons", "hackathons"), ("Workshops", "workshops"),
            ("Competitions", "competitions"),
        ],
        "Interests": [
            ("Hobbies", "hobbies"), ("Interest area", "interestArea"), ("Career fields", "careerFields"),
        ],
        "Future Goals": [
            ("Career", "career"), ("Skill to learn", "goalSkill"), ("Two-year goal", "twoYearGoal"),
            ("Organization", "organization"), ("Self-improvement", "selfImprovement"),
        ],
    }
    for sheet_name, fields in layouts.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["Student profile ID", *(label for label, _ in fields)])
        worksheet.append([profile_id, *(excel_value(profile.get(key)) for _, key in fields)])
        style_worksheet(worksheet)

    for sheet_name, record_key, fields in (
        ("Certifications", "certifications", [("Name", "name"), ("Organization", "organization"), ("Year", "year")]),
        ("Achievements", "achievements", [("Title", "title"), ("Type", "type"), ("Year", "year"), ("Description", "description")]),
    ):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["Student profile ID", *(label for label, _ in fields)])
        for record in profile.get(record_key, []):
            worksheet.append([profile_id, *(excel_value(record.get(key)) for _, key in fields)])
        style_worksheet(worksheet)
    return workbook


def generate_profile_workbook(account_id: int) -> tuple[Path, int]:
    with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            profile_id, profile = read_student_profile(cursor, account_id)
    if profile_id is None or profile is None:
        raise LookupError("No saved profile exists for this account.")

    EXPORT_DIR.mkdir(exist_ok=True)
    target = EXPORT_DIR / f"student_profile_{profile_id}.xlsx"
    temporary = EXPORT_DIR / f"student_profile_{profile_id}.{secrets.token_hex(8)}.tmp.xlsx"
    save_workbook_atomically(build_profile_workbook(profile_id, profile), target, temporary)
    return target, profile_id


def generate_all_profiles_workbook() -> Path:
    with psycopg.connect(get_database_url(), connect_timeout=10) as connection:
        with connection.cursor(row_factory=dict_row) as cursor:
            cursor.execute("SELECT * FROM student_profile_overview ORDER BY id")
            profiles = [dict(row) for row in cursor.fetchall()]

    workbook = Workbook()
    workbook.remove(workbook.active)
    layouts = {
        "Personal": [
            ("Full name", "full_name"), ("Preferred name", "preferred_name"), ("Age", "age"),
            ("Date of birth", "date_of_birth"), ("Gender", "gender"), ("Phone", "phone"),
            ("City", "city"), ("Created at (UTC)", "created_at"),
        ],
        "Education": [
            ("College / school", "school"), ("Course", "course"), ("Department", "department"),
            ("Year", "current_year"), ("Semester", "current_semester"),
            ("Student ID", "campus_student_id"), ("CGPA / percentage", "grade"),
        ],
        "Academics": [
            ("Favorite subjects", "favorite_subjects"), ("Difficult subject", "difficult_subject"),
            ("Studying subjects", "studying_subjects"), ("Study hours", "study_hours"),
            ("Learning methods", "learning_methods"), ("Academic confidence", "academic_confidence"),
        ],
        "Technical Skills": [
            ("Programming languages", "programming_languages"), ("Technical skills", "technical_skills"),
            ("Strongest skill", "strongest_skill"), ("Next technology", "next_technology"),
            ("Skill level", "skill_level"),
        ],
        "Activities": [
            ("Clubs", "clubs"), ("Hackathons", "hackathons"),
            ("Workshops", "workshops"), ("Competitions", "competitions"),
        ],
        "Interests": [
            ("Hobbies", "hobbies"), ("Interest area", "interest_area"), ("Career fields", "career_fields"),
        ],
        "Future Goals": [
            ("Career", "career"), ("Skill to learn", "goal_skill"), ("Two-year goal", "two_year_goal"),
            ("Organization", "organization"), ("Self-improvement", "self_improvement"),
        ],
    }
    for sheet_name, fields in layouts.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["Student profile ID", "Email", *(label for label, _ in fields)])
        for profile in profiles:
            worksheet.append([
                profile["id"], profile.get("email"),
                *(excel_value(profile.get(key)) for _, key in fields),
            ])
        style_worksheet(worksheet)

    for sheet_name, record_key, fields in (
        ("Certifications", "certifications", [("Name", "name"), ("Organization", "organization"), ("Year", "year")]),
        ("Achievements", "achievements", [("Title", "title"), ("Type", "type"), ("Year", "year"), ("Description", "description")]),
    ):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["Student profile ID", "Email", "Full name", *(label for label, _ in fields)])
        for profile in profiles:
            for record in profile.get(record_key) or []:
                worksheet.append([
                    profile["id"], profile.get("email"), profile.get("full_name"),
                    *(excel_value(record.get(key)) for _, key in fields),
                ])
        style_worksheet(worksheet)

    EXPORT_DIR.mkdir(exist_ok=True)
    target = EXPORT_DIR / "all_student_profiles.xlsx"
    temporary = EXPORT_DIR / f"all_student_profiles.{secrets.token_hex(8)}.tmp.xlsx"
    save_workbook_atomically(workbook, target, temporary)
    return target


def yes_no_to_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    return value == "Yes"


def record_has_values(record: BaseModel) -> bool:
    return any(value is not None for value in record.model_dump().values())


def save_profile_sections(cursor: psycopg.Cursor, profile: ProfileSubmission, student_id: int) -> None:
    cursor.execute(
        """
        INSERT INTO education
            (student_id, school, course, department, current_year, current_semester, campus_student_id, grade)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            school = EXCLUDED.school, course = EXCLUDED.course, department = EXCLUDED.department,
            current_year = EXCLUDED.current_year, current_semester = EXCLUDED.current_semester,
            campus_student_id = EXCLUDED.campus_student_id, grade = EXCLUDED.grade
        """,
        (student_id, profile.school, profile.course, profile.department, profile.year, profile.semester, profile.student_id, profile.grade),
    )
    cursor.execute(
        """
        INSERT INTO academic_profile
            (student_id, favorite_subjects, difficult_subject, studying_subjects, study_hours,
             learning_methods, academic_confidence)
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            favorite_subjects = EXCLUDED.favorite_subjects, difficult_subject = EXCLUDED.difficult_subject,
            studying_subjects = EXCLUDED.studying_subjects, study_hours = EXCLUDED.study_hours,
            learning_methods = EXCLUDED.learning_methods, academic_confidence = EXCLUDED.academic_confidence
        """,
        (student_id, profile.favorite_subjects, profile.difficult_subject, profile.studying_subjects, profile.study_hours, profile.learning_methods, profile.academic_confidence),
    )
    cursor.execute(
        """
        INSERT INTO technical_profile
            (student_id, programming_languages, technical_skills, strongest_skill, next_technology, skill_level)
        VALUES (%s, %s, %s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            programming_languages = EXCLUDED.programming_languages, technical_skills = EXCLUDED.technical_skills,
            strongest_skill = EXCLUDED.strongest_skill, next_technology = EXCLUDED.next_technology,
            skill_level = EXCLUDED.skill_level
        """,
        (student_id, profile.languages, profile.technical_skills, profile.strongest_skill, profile.next_technology, profile.technical_level),
    )
    cursor.execute(
        """
        INSERT INTO activities (student_id, clubs, hackathons, workshops, competitions)
        VALUES (%s, %s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            clubs = EXCLUDED.clubs, hackathons = EXCLUDED.hackathons,
            workshops = EXCLUDED.workshops, competitions = EXCLUDED.competitions
        """,
        (student_id, profile.clubs, yes_no_to_bool(profile.hackathons), yes_no_to_bool(profile.workshops), yes_no_to_bool(profile.competitions)),
    )

    cursor.execute("DELETE FROM certifications WHERE student_id = %s", (student_id,))
    for certification in profile.certifications:
        if record_has_values(certification):
            cursor.execute(
                "INSERT INTO certifications (student_id, name, organization, year) VALUES (%s, %s, %s, %s)",
                (student_id, certification.name, certification.organization, certification.year),
            )

    cursor.execute("DELETE FROM achievements WHERE student_id = %s", (student_id,))
    for achievement in profile.achievements:
        if record_has_values(achievement):
            cursor.execute(
                "INSERT INTO achievements (student_id, title, type, year, description) VALUES (%s, %s, %s, %s, %s)",
                (student_id, achievement.title, achievement.type, achievement.year, achievement.description),
            )

    cursor.execute(
        """
        INSERT INTO interests (student_id, hobbies, interest_area, career_fields)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            hobbies = EXCLUDED.hobbies, interest_area = EXCLUDED.interest_area,
            career_fields = EXCLUDED.career_fields
        """,
        (student_id, profile.hobbies, profile.interest_area, profile.career_fields),
    )
    cursor.execute(
        """
        INSERT INTO future_goals (student_id, career, goal_skill, two_year_goal, organization, self_improvement)
        VALUES (%s, %s, %s, %s, %s, %s)
        ON CONFLICT (student_id) DO UPDATE SET
            career = EXCLUDED.career, goal_skill = EXCLUDED.goal_skill,
            two_year_goal = EXCLUDED.two_year_goal, organization = EXCLUDED.organization,
            self_improvement = EXCLUDED.self_improvement
        """,
        (student_id, profile.career, profile.goal_skill, profile.two_year_goal, profile.organization, profile.self_improvement),
    )


app = FastAPI(title="Student Profile Portal", docs_url=None, redoc_url=None)


@app.get("/api/health")
def database_health() -> dict[str, str]:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=3) as connection:
            connection.execute("SELECT 1")
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable. Check DATABASE_URL.") from None
    return {"status": "connected"}


@app.post("/api/register", status_code=201)
def register(credentials: Credentials, response: Response) -> dict[str, object]:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute(
                    "INSERT INTO accounts (email, password_hash) VALUES (%s, %s) RETURNING id, email",
                    (credentials.email, hash_password(credentials.password)),
                )
                account = cursor.fetchone()
                cursor.execute(
                    """
                    SELECT id FROM students
                    WHERE account_id IS NULL AND lower(email) = %s
                    ORDER BY created_at DESC, id DESC LIMIT 1 FOR UPDATE
                    """,
                    (credentials.email,),
                )
                legacy_profile = cursor.fetchone()
                if legacy_profile:
                    cursor.execute(
                        "UPDATE students SET account_id = %s WHERE id = %s",
                        (account["id"], legacy_profile["id"]),
                    )
                token = create_session(cursor, account["id"])
                profile_id, profile = read_student_profile(cursor, account["id"])
    except psycopg.errors.UniqueViolation:
        raise HTTPException(status_code=409, detail="An account with this email already exists. Log in instead.") from None
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable.") from None
    except psycopg.errors.UndefinedTable:
        raise HTTPException(status_code=503, detail="Rerun the updated schema.sql in pgAdmin.") from None
    set_session_cookie(response, token)
    return {"id": account["id"], "email": account["email"], "profileId": profile_id, "profile": profile}


@app.post("/api/login")
def login(credentials: Credentials, response: Response) -> dict[str, object]:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                cursor.execute("SELECT id, email, password_hash FROM accounts WHERE email = %s", (credentials.email,))
                account = cursor.fetchone()
                if account is None or not verify_password(credentials.password, account["password_hash"]):
                    raise HTTPException(status_code=401, detail="Email or password is incorrect.")
                token = create_session(cursor, account["id"])
                profile_id, profile = read_student_profile(cursor, account["id"])
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable.") from None
    except psycopg.errors.UndefinedTable:
        raise HTTPException(status_code=503, detail="Rerun the updated schema.sql in pgAdmin.") from None
    set_session_cookie(response, token)
    return {"id": account["id"], "email": account["email"], "profileId": profile_id, "profile": profile}


@app.get("/api/me")
def current_account(request: Request) -> dict[str, object]:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                account = authenticated_account(cursor, request)
                profile_id, profile = read_student_profile(cursor, account["id"])
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable.") from None
    except psycopg.errors.UndefinedTable:
        raise HTTPException(status_code=503, detail="Rerun the updated schema.sql in pgAdmin.") from None
    return {"id": account["id"], "email": account["email"], "profileId": profile_id, "profile": profile}


@app.get("/api/profiles/excel")
def download_profile_excel(request: Request) -> FileResponse:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                account = authenticated_account(cursor, request)
                profile_id, _ = read_student_profile(cursor, account["id"])
        if profile_id is None:
            raise HTTPException(status_code=404, detail="Complete your profile before downloading the Excel workbook.")
        workbook_path, profile_id = generate_profile_workbook(int(account["id"]))
    except HTTPException:
        raise
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable.") from None
    except psycopg.errors.UndefinedTable:
        raise HTTPException(status_code=503, detail="Rerun the updated schema.sql in pgAdmin.") from None
    except Exception as error:
        logger.error("Excel workbook generation failed: %s", type(error).__name__)
        raise HTTPException(status_code=500, detail="The Excel workbook could not be generated.") from None
    return FileResponse(
        workbook_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"student_profile_{profile_id}.xlsx",
    )


@app.post("/api/logout")
def logout(request: Request, response: Response) -> dict[str, bool]:
    token = request.cookies.get(SESSION_COOKIE)
    if token and DATABASE_URL:
        token_hash = hashlib.sha256(token.encode()).hexdigest()
        try:
            with psycopg.connect(DATABASE_URL, connect_timeout=3) as connection:
                connection.execute("DELETE FROM auth_sessions WHERE token_hash = %s", (token_hash,))
        except psycopg.Error:
            logger.warning("Could not remove login session during logout")
    response.delete_cookie(SESSION_COOKIE, path="/", httponly=True, samesite="lax")
    return {"loggedOut": True}


@app.put("/api/profiles")
def save_profile(profile: ProfileSubmission, request: Request) -> dict[str, int | bool]:
    try:
        with psycopg.connect(get_database_url(), connect_timeout=5) as connection:
            with connection.cursor(row_factory=dict_row) as cursor:
                account = authenticated_account(cursor, request)
                account_id = int(account["id"])
                cursor.execute("SELECT id FROM accounts WHERE id = %s FOR UPDATE", (account_id,))
                cursor.execute("SELECT id FROM students WHERE account_id = %s FOR UPDATE", (account_id,))
                existing = cursor.fetchone()
                if existing:
                    student_id = existing["id"]
                    cursor.execute(
                        """
                        UPDATE students SET full_name = %s, preferred_name = %s, age = %s,
                            date_of_birth = %s, gender = %s, email = %s, phone = %s, city = %s
                        WHERE id = %s
                        """,
                        (profile.full_name, profile.preferred_name, profile.age, profile.date_of_birth,
                         profile.gender, account["email"], profile.phone, profile.city, student_id),
                    )
                else:
                    cursor.execute(
                        """
                        INSERT INTO students
                            (account_id, full_name, preferred_name, age, date_of_birth,
                             gender, email, phone, city)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                        RETURNING id
                        """,
                        (account_id, profile.full_name, profile.preferred_name, profile.age, profile.date_of_birth,
                         profile.gender, account["email"], profile.phone, profile.city),
                    )
                    student_id = cursor.fetchone()["id"]
                save_profile_sections(cursor, profile, student_id)
    except HTTPException:
        raise
    except psycopg.OperationalError:
        raise HTTPException(status_code=503, detail="PostgreSQL is unreachable. Check DATABASE_URL and retry.") from None
    except psycopg.errors.UndefinedTable:
        raise HTTPException(status_code=503, detail="Rerun the updated schema.sql in pgAdmin.") from None
    except psycopg.Error as error:
        logger.error("PostgreSQL profile save failed: %s", type(error).__name__)
        raise HTTPException(status_code=500, detail="The profile could not be saved. Check the schema and retry.") from None
    excel_saved = False
    try:
        generate_profile_workbook(account_id)
        excel_saved = True
    except Exception as error:
        logger.error("Excel workbook generation failed: %s", type(error).__name__)
    all_excel_saved = False
    try:
        generate_all_profiles_workbook()
        all_excel_saved = True
    except Exception as error:
        logger.error("All-student Excel workbook generation failed: %s", type(error).__name__)
    return {
        "id": student_id,
        "saved": True,
        "excelSaved": excel_saved,
        "allExcelSaved": all_excel_saved,
    }


@app.get("/", include_in_schema=False)
def home() -> FileResponse:
    return FileResponse(ROOT / "index.html")


@app.get("/styles.css", include_in_schema=False)
def styles() -> FileResponse:
    return FileResponse(ROOT / "styles.css", media_type="text/css")


@app.get("/script.js", include_in_schema=False)
def script() -> FileResponse:
    return FileResponse(ROOT / "script.js", media_type="text/javascript")


if __name__ == "__main__":
    uvicorn.run(app, host="127.0.0.1", port=8000)
